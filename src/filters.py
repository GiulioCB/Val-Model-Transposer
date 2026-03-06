from __future__ import annotations

from typing import List, Dict, Any, Tuple
from openpyxl import load_workbook
from .excel_utils import col_to_idx, normalize_str, is_blank
from .types import FilterQuestion

DROPDOWN_QS = set(range(1, 8))        # 1..7
TEXT_QS = {8, 9, 15, 16}              # free input
YESNO_QS = set(range(10, 15))         # 10..14

def read_filter_questions(
    output_template_path: str,
    filters_sheet_name: str,
    question_row: int,
    required_row: int,
    start_col: str,
    end_col: str,
    dropdown_start_row: int,
) -> List[FilterQuestion]:
    wb = load_workbook(output_template_path, data_only=True)
    if filters_sheet_name not in wb.sheetnames:
        raise ValueError(f"Template missing sheet: {filters_sheet_name}")
    ws = wb[filters_sheet_name]

    s = col_to_idx(start_col)
    e = col_to_idx(end_col)

    questions: List[FilterQuestion] = []
    for i, col_idx in enumerate(range(s, e + 1), start=1):
        title = normalize_str(ws.cell(row=question_row, column=col_idx).value)
        required_flag = normalize_str(ws.cell(row=required_row, column=col_idx).value).lower()
        required = required_flag == "x"

        if i in DROPDOWN_QS:
            opts = []
            r = dropdown_start_row
            while True:
                v = ws.cell(row=r, column=col_idx).value
                if is_blank(v):
                    break
                opts.append(normalize_str(v))
                r += 1
            kind = "dropdown"
            questions.append(FilterQuestion(idx=i, title=title or f"Filter {i}", required=required, kind=kind, options=opts))
        elif i in TEXT_QS:
            questions.append(FilterQuestion(idx=i, title=title or f"Filter {i}", required=required, kind="text", options=None))
        elif i in YESNO_QS:
            questions.append(FilterQuestion(idx=i, title=title or f"Filter {i}", required=required, kind="yesno", options=["Yes", "No"]))
        else:
            # fallback: treat as text
            questions.append(FilterQuestion(idx=i, title=title or f"Filter {i}", required=required, kind="text", options=None))

    return questions

def validate_answers(questions: List[FilterQuestion], answers: Dict[int, Any]) -> Tuple[bool, List[str]]:
    missing: List[str] = []
    for q in questions:
        if not q.required:
            continue
        v = answers.get(q.idx, None)
        if q.kind in ("dropdown", "text"):
            if v is None or normalize_str(v) == "":
                missing.append(f"Filter {q.idx}: {q.title}")
        elif q.kind == "yesno":
            if v not in ("Yes", "No"):
                missing.append(f"Filter {q.idx}: {q.title}")
        else:
            if v is None or normalize_str(v) == "":
                missing.append(f"Filter {q.idx}: {q.title}")
    return (len(missing) == 0, missing)
