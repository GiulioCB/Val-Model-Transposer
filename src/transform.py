from __future__ import annotations

from typing import Any, Dict, List, Tuple
from zipfile import BadZipFile
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import column_index_from_string
from .excel_utils import col_to_idx, normalize_str
from .geocode import GeocodeConfig, build_address, geocode_address

def _get_ws(wb, name: str) -> Worksheet:
    if name not in wb.sheetnames:
        raise ValueError(f"Input workbook missing sheet: {name}")
    return wb[name]

def read_start_page_fields(ws_start) -> Dict[str, Any]:
    # Direct cell reads requested by user
    return {
        "property_name": ws_start["C4"].value,
        "street_address": ws_start["C5"].value,
        "country_region": ws_start["C8"].value,
        "city_town": ws_start["C6"].value,
        "zip_post_code": ws_start["C7"].value,
        "rooms": ws_start["C11"].value,
        "classification": ws_start["F10"].value,
        "meeting_rooms": ws_start["F15"].value,
        "meeting_max_capacity": ws_start["F16"].value,
        "currency": ws_start["C21"].value,
        "pcd2": ws_start["C7"].value,  # as specified
    }

def find_qualifying_projection_columns(
    ws_proj,
    scan_start_col: str,
    scan_end_col: str,
    last_month_row: int,
    status_row: int,
) -> List[int]:
    s = col_to_idx(scan_start_col)
    e = col_to_idx(scan_end_col)
    qualifying: List[int] = []
    for col_idx in range(s, e + 1):
        last_month = normalize_str(ws_proj.cell(row=last_month_row, column=col_idx).value)
        status = normalize_str(ws_proj.cell(row=status_row, column=col_idx).value)
        if last_month.lower() == "december" and status.lower() == "actual":
            qualifying.append(col_idx)
    return qualifying

def extract_transpose_vector(
    ws_proj,
    col_idx: int,
    start_row: int,
    end_row: int
) -> List[Any]:
    return [ws_proj.cell(row=r, column=col_idx).value for r in range(start_row, end_row + 1)]


def read_projection_row_headers(
    ws_proj,
    start_row: int,
    end_row: int,
    title_col: int = 2,
) -> List[Any]:
    return [ws_proj.cell(row=r, column=title_col).value for r in range(start_row, end_row + 1)]

def build_static_row_values(
    start_fields: Dict[str, Any],
    filter_answers: Dict[int, Any],
    geocode_cfg: GeocodeConfig
) -> Tuple[List[Any], List[str]]:
    warnings: List[str] = []

    # Build address for geocode from C5, C7, C6, C8
    street = normalize_str(start_fields.get("street_address"))
    zip_code = normalize_str(start_fields.get("zip_post_code"))
    city = normalize_str(start_fields.get("city_town"))
    country = normalize_str(start_fields.get("country_region"))
    address = build_address(street=street, zip_code=zip_code, city=city, country=country)

    lat, lon, geo_warn = geocode_address(address, geocode_cfg)
    if geo_warn:
        warnings.append(geo_warn)

    # Map filters by number
    f = filter_answers
    # Yes/No filters (10-14) are stored as "Yes"/"No"
    static_values = [
        start_fields.get("property_name"),                 # A
        None,                                              # B StreetNumber
        start_fields.get("street_address"),                # C StreetAddress
        None,                                              # D StreetType
        None,                                              # E StreetPrefix
        None,                                              # F StreetSuffix
        None,                                              # G POBox
        None,                                              # H Country
        start_fields.get("country_region"),                # I Country - Region
        None,                                              # J Region
        None,                                              # K County
        start_fields.get("city_town"),                     # L City/Town
        None,                                              # M District
        start_fields.get("zip_post_code"),                 # N Zip/Post Code
        f.get(1),                                          # O PropertyType
        None,                                              # P Owner Type
        f.get(2),                                          # Q Location
        f.get(3),                                          # R Food Bev Operator
        f.get(4),                                          # S Operator
        start_fields.get("rooms"),                         # T Rooms
        None,                                              # U Metro Area
        None,                                              # V Market Area
        None,                                              # W Submarket Area
        f.get(5),                                          # X Chain/ChainID
        start_fields.get("classification"),                # Y Classification
        f.get(6),                                          # Z Management Company
        f.get(7),                                          # AA Owner Company
        f.get(8),                                          # AB YearOpened
        None,                                              # AC YearClosed
        None,                                              # AD Year Recent Renovation
        None,                                              # AE OwnBuildings
        None,                                              # AF OwnLand
        f.get(9),                                          # AG MeetingSpace (SQM)
        start_fields.get("meeting_rooms"),                 # AH MeetingRooms
        start_fields.get("meeting_max_capacity"),          # AI MeetingMaxCapacity
        None,                                              # AJ Casino
        None,                                              # AK Convention
        None,                                              # AL Conference
        f.get(10),                                         # AM Ski
        f.get(11),                                         # AN Spa
        f.get(12),                                         # AO HealthClub
        f.get(13),                                         # AP Golf
        f.get(14),                                         # AQ Boutique
        None,                                              # AR AllSuite
        None,                                              # AS Suites
        None,                                              # AT Floors
        None,                                              # AU ParkingSpaces
        f.get(15),                                         # AV FoodOutlets
        f.get(16),                                         # AW BeverageOutlets
        None,                                              # AX Financial Data provider
        None,                                              # AY Notes
        start_fields.get("pcd2"),                          # AZ pcd2
        lat,                                               # BA lat (Automation 1)
        lon,                                               # BB long (Automation 2)
        start_fields.get("currency"),                      # BC Currency
    ]
    return static_values, warnings

def load_input_workbook(path: str):
    # data_only=True reads cached results (Excel-calculated values). See README notes.
    try:
        return load_workbook(path, data_only=True, keep_vba=False)
    except BadZipFile as exc:
        raise ValueError(
            "The uploaded input workbook is not a valid Excel .xlsx/.xlsm file. "
            "Please re-save the source model in Excel and upload it again."
        ) from exc
