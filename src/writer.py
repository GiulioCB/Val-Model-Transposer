from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any, List
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook

from .excel_utils import col_to_idx
from .github_store import download_file_bytes, upload_file_bytes


@dataclass
class DestinationWorkbookSession:
    workbook: Any
    destination_path: str | None = None
    github_repo: str | None = None
    github_excel_path: str | None = None
    github_branch: str | None = None
    github_sha: str | None = None
    keep_vba: bool = False


def _is_macro_enabled_workbook(path: str | None) -> bool:
    if not path:
        return False
    return Path(path).suffix.lower() == ".xlsm"


def open_destination_workbook(
    destination_path: str | None = None,
    github_repo: str | None = None,
    github_excel_path: str | None = None,
    github_branch: str | None = None,
) -> DestinationWorkbookSession:
    keep_vba = _is_macro_enabled_workbook(github_excel_path or destination_path)

    if github_repo and github_excel_path:
        workbook_bytes, github_sha = download_file_bytes(github_repo, github_excel_path, github_branch)
        try:
            workbook = load_workbook(BytesIO(workbook_bytes), keep_vba=keep_vba)
        except BadZipFile as exc:
            raise ValueError(
                f"The destination workbook '{github_excel_path}' in GitHub is not a valid Excel .xlsx/.xlsm file. "
                "Upload a clean Excel workbook to the repo and try again."
            ) from exc
        return DestinationWorkbookSession(
            workbook=workbook,
            github_repo=github_repo,
            github_excel_path=github_excel_path,
            github_branch=github_branch,
            github_sha=github_sha,
            keep_vba=keep_vba,
        )

    if not destination_path:
        raise ValueError("A local destination path or GitHub workbook location is required.")

    try:
        workbook = load_workbook(destination_path, keep_vba=keep_vba)
    except BadZipFile as exc:
        raise ValueError(
            f"The destination workbook '{destination_path}' is not a valid Excel .xlsx/.xlsm file."
        ) from exc
    return DestinationWorkbookSession(workbook=workbook, destination_path=destination_path, keep_vba=keep_vba)


def insert_and_write_row_in_workbook(
    workbook: Any,
    sheet_name: str,
    insert_row_index: int,
    static_values_A_to_BC: List[Any],
    dynamic_values_BD_to_LZ: List[Any],
    static_start_col: str = "A",
    dynamic_start_col: str = "BD",
) -> None:
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Destination workbook missing sheet: {sheet_name}")
    ws = workbook[sheet_name]

    ws.insert_rows(insert_row_index, amount=1)

    start_idx = col_to_idx(static_start_col)
    for offset, val in enumerate(static_values_A_to_BC):
        ws.cell(row=insert_row_index, column=start_idx + offset).value = val

    dyn_start_idx = col_to_idx(dynamic_start_col)
    for offset, val in enumerate(dynamic_values_BD_to_LZ):
        ws.cell(row=insert_row_index, column=dyn_start_idx + offset).value = val


def read_header_row(workbook: Any, sheet_name: str, end_col_idx: int) -> List[Any]:
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Destination workbook missing sheet: {sheet_name}")
    ws = workbook[sheet_name]
    return [ws.cell(row=1, column=col_idx).value for col_idx in range(1, end_col_idx + 1)]


def save_destination_workbook(session: DestinationWorkbookSession, commit_message: str | None = None) -> None:
    workbook_to_save = session.workbook

    if not session.keep_vba:
        clean_workbook = Workbook()
        clean_workbook.remove(clean_workbook.active)

        for sheet_name in session.workbook.sheetnames:
            source_ws = session.workbook[sheet_name]
            target_ws = clean_workbook.create_sheet(title=sheet_name)
            for row in source_ws.iter_rows():
                for cell in row:
                    target_ws[cell.coordinate].value = cell.value

        workbook_to_save = clean_workbook

    if session.github_repo and session.github_excel_path:
        output = BytesIO()
        workbook_to_save.save(output)
        upload_file_bytes(
            repo_name=session.github_repo,
            path=session.github_excel_path,
            content=output.getvalue(),
            sha=session.github_sha,
            message=commit_message or "Update OUTPUT.xlsx via Streamlit",
            branch=session.github_branch,
        )
        return

    if not session.destination_path:
        raise ValueError("Destination path is required for local workbook saves.")

    workbook_to_save.save(session.destination_path)
